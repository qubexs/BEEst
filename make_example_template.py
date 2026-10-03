"""Create templates/example_BE2026.xlsx so the app works out of the box.

Replace this file with YOUR real form later - the app auto-detects any
A=Indicator / B=Valoare layout or {{PLACEHOLDER}} tokens.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

rows = [
    ("FORMULAR BILANT ECONOMIC 2026 (exemplu - inlocuiti cu formularul real)", ""),
    ("Indicator", "Valoare"),
    ("Denumire societate", ""),
    ("CUI / Cod fiscal", ""),
    ("Nr. Reg. Com.", ""),
    ("Adresa sediu", ""),
    ("Cod CAEN", ""),
    ("Telefon", ""),
    ("An fiscal", ""),
    ("Curs EUR/RON (BNR)", ""),
    ("Curs USD/RON (BNR)", ""),
    ("Inflatie anuala IPC % (INSSE)", ""),
    ("Crestere PIB % (estimare)", ""),
    ("Cifra de afaceri neta (lei)", ""),
    ("Profit net (lei)", ""),
    ("Numar salariati", ""),
    ("Observatii", ""),
]

out = Path(__file__).resolve().parent / "templates" / "example_BE2026.xlsx"
out.parent.mkdir(exist_ok=True)
wb = Workbook()
ws = wb.active
ws.title = "BE2026"
ws.sheet_properties.pageSetUpPr.fitToPage = True
thin = Side(style="thin", color="999999")
for i, (a, b) in enumerate(rows, start=1):
    ws.cell(i, 1, a)
    ws.cell(i, 2, b)
    for col in (1, 2):
        c = ws.cell(i, 1 if col == 1 else 2)
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    if i == 1:
        ws.cell(i, 1).font = Font(bold=True, size=13)
        ws.cell(i, 1).fill = PatternFill("solid", fgColor="D9E2F3")
        ws.merge_cells("A1:B1")
        ws.row_dimensions[1].height = 30
    elif i == 2:
        ws.cell(i, 1).font = Font(bold=True)
        ws.cell(i, 2).font = Font(bold=True)
        ws.cell(i, 1).fill = PatternFill("solid", fgColor="E2EFDA")
        ws.cell(i, 2).fill = PatternFill("solid", fgColor="E2EFDA")
ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 38
wb.save(out)
print(f"OK: {out}")
