"""Read any BE2026 xlsx template and detect fillable fields — fast streaming scan.

Layouts (auto-detected):
1. Two-column form:  A=Label, B=Value (empty)  -> fillable cell B
2. {{PLACEHOLDER}} tokens in visible cells       -> fillable cell
3. Table questionnaire (KP205-style): coded rows (A/B=code, B/C=item) with
   input columns; header cover-block on top    -> fillable empty cells

Speed rules: read_only streaming, hidden sheets skipped, single row pass +
table pass over cached grid, placeholder width capped, empty-streak break,
result cached by (path, size, mtime).
"""
from __future__ import annotations
from dataclasses import dataclass
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import os
import re

PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z0-9_\-.\s/]+?)\s*\}\}")
CODE_RE = re.compile(r"^\d[\d.\s()a-z\-]{0,13}$", re.I)

MAX_PH_COLS = 60        # placeholder scan width cap
MAX_TABLE_COLS = 60     # table scan width cap (cols D..)
MAX_EMPTY_STREAK = 500  # stop sheet after this many fully-empty rows
MAX_FIELDS_SHEET = 5000  # safety cap per sheet
HEADER_LABELS = {"indicator", "indicatorul", "denumire indicator",
                 "label", "field", "penunjuk"}

_cache: dict[tuple, tuple[list, list]] = {}


@dataclass
class TemplateField:
    sheet: str
    cell: str          # e.g. "B5"
    label: str         # human label
    current: object    # current cell value
    kind: str          # "key_value" | "placeholder" | "table_cell"


def _cache_key(path: str):
    try:
        st = os.stat(path)
        return (os.path.abspath(path), st.st_size, st.st_mtime_ns)
    except OSError:
        return None


def _is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _is_code(v) -> bool:
    return isinstance(v, str) and bool(CODE_RE.match(v.strip()))


def _text(v) -> str:
    return v.strip() if isinstance(v, str) and not _is_formula(v) else ""


def inspect_template(path: str) -> tuple[list[TemplateField], list[str]]:
    key = _cache_key(path)
    if key is not None and key in _cache:
        return _cache[key]

    wb = load_workbook(path, read_only=True)
    try:
        sheets = [ws.title for ws in wb.worksheets]
        fields: list[TemplateField] = []
        seen: set[tuple[str, str]] = set()

        def add(sheet: str, cell: str, label: str, current: object, kind: str):
            if (sheet, cell) not in seen:
                seen.add((sheet, cell))
                fields.append(TemplateField(sheet, cell, label, current, kind))

        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            grid = [[c.value for c in row] for row in ws.iter_rows()]
            nrows = len(grid)
            if not nrows:
                continue

            def val(r: int, c: int):
                row = grid[r]
                return row[c] if c < len(row) else None

            # ---- pass 1+2: placeholders & A/B key-value (as before) ----
            empty_streak = 0
            for r in range(nrows):
                row_empty = True
                width = min(len(grid[r]), MAX_PH_COLS)
                for cc in range(width):
                    v = grid[r][cc]
                    if v is None or v == "":
                        continue
                    row_empty = False
                    if isinstance(v, str):
                        for m in PLACEHOLDER_RE.finditer(v):
                            lab = m.group(1).strip()
                            if lab:
                                add(ws.title, f"{get_column_letter(cc + 1)}{r + 1}",
                                    lab, v, "placeholder")
                a = grid[r][0] if len(grid[r]) > 0 else None
                b = grid[r][1] if len(grid[r]) > 1 else None
                ta, tb = _text(a), _text(b)
                if len(ta) >= 2:
                    row_empty = False
                    b_empty = (b is None or b == "" or isinstance(b, (int, float)))
                    has_ph = isinstance(b, str) and "{{" in b
                    is_title = len(ta) > 60 and not has_ph and b in (None, "")
                    is_formula_a = isinstance(a, str) and a.startswith("=")
                    if (b_empty or has_ph) and not is_title \
                            and not is_formula_a and ta.lower() not in HEADER_LABELS:
                        add(ws.title, f"B{r + 1}", ta, b, "key_value")
                elif b not in (None, ""):
                    row_empty = False
                if row_empty:
                    empty_streak += 1
                    if empty_streak >= MAX_EMPTY_STREAK:
                        break
                else:
                    empty_streak = 0

            # ---- pass 3: table questionnaire ----
            _detect_table(ws.title, grid, add)
        result = (fields, sheets)
        if key is not None:
            _cache[key] = result
        return result
    finally:
        wb.close()


def _detect_table(sheet: str, grid: list, add) -> None:
    """Coded-row tables: empty non-formula cells in live columns become fields."""
    nrows = len(grid)
    width = min(max((len(r) for r in grid), default=0), MAX_TABLE_COLS)
    if width <= 3 or nrows < 3:
        return
    made = [0]

    def put(cell: str, label: str, v) -> bool:
        if made[0] >= MAX_FIELDS_SHEET:
            return False
        add(sheet, cell, label, v, "table_cell")
        made[0] += 1
        return True

    def val(r: int, c: int):
        row = grid[r]
        return row[c] if c < len(row) else None

    def row_has_formula(r: int, c0: int = 3) -> bool:
        return any(_is_formula(val(r, c)) for c in range(c0, min(len(grid[r]), width)))

    def row_empty(r: int) -> bool:
        return all(val(r, c) in (None, "") for c in range(min(len(grid[r]), width)))

    # header zone end: first coded row (A or B code + item text)
    r0 = None
    for r in range(nrows):
        a, b, c = _text(val(r, 0)), _text(val(r, 1)), _text(val(r, 2))
        if (_is_code(a) and (b or c)) or (_is_code(b) and c):
            r0 = r
            break
    if r0 is None:
        r0 = nrows  # no coded rows: whole sheet is cover-style

    def live_cols(lo: int, hi: int) -> set[int]:
        live = set()
        for cc in range(3, width):
            for r in range(lo, min(hi, nrows)):
                v = val(r, cc)
                if v is not None and v != "":
                    live.add(cc)
                    break
        return live

    def colhead(r: int, cc: int) -> str:
        for rr in range(r - 1, max(-1, r - 9), -1):
            v = val(rr, cc)
            if isinstance(v, str) and v.strip() and not _is_formula(v):
                return v.strip()[:40]
        return get_column_letter(cc + 1)

    # --- zone A: cover block (rows before first coded row) ---
    if r0 > 0:
        live = live_cols(0, r0)
        for r in range(r0):
            b = _text(val(r, 1))
            if len(b) < 2 or not row_has_formula(r):
                continue
            for cc in sorted(live):
                v = val(r, cc)
                if v is None or v == "":
                    ch = colhead(r, cc)
                    lab = b if ch == get_column_letter(cc + 1) else f"{b} › {ch}"
                    if not put(f"{get_column_letter(cc + 1)}{r + 1}", lab, v):
                        return

    # --- zone B: table blocks split by fully-empty rows ---
    r = r0
    while r < nrows:
        if row_empty(r):
            r += 1
            continue
        lo = r
        while r < nrows and not row_empty(r):
            r += 1
        hi = r  # block [lo, hi)
        live = live_cols(lo, hi)
        if not live:
            continue
        section = ""
        for rr in range(lo, hi):
            a, b, c = _text(val(rr, 0)), _text(val(rr, 1)), _text(val(rr, 2))
            texts_d = sum(1 for cc in range(3, width)
                          if isinstance(val(rr, cc), str) and val(rr, cc).strip()
                          and not _is_formula(val(rr, cc)))
            if texts_d >= 3 and not (_is_code(a) or _is_code(b)):
                continue  # column-header row: skip as field, usable as colhead
            code = None
            item = ""
            if _is_code(a) and (b or c):
                code, item = a, c or b
            elif _is_code(b) and c:
                code, item = b, c
            if code and item:
                for cc in sorted(live):
                    v = val(rr, cc)
                    if v is None or v == "":
                        ch = colhead(rr, cc)
                        parts = ([section] if section else []) + [f"{code} {item}", ch]
                        if not put(f"{get_column_letter(cc + 1)}{rr + 1}",
                                   " › ".join(parts), v):
                            return
            elif b and not c and not _is_code(b):
                section = b  # section header row
